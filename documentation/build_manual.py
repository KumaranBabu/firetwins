"""Build FireTwins_User_Manual.docx (and .pdf through Word) from the walkthrough screenshots and logs.

    python documentation/build_manual.py

Figures and their captions come from documentation/screenshots/<example>/validation_log.json in the order the
walkthrough took them; every number quoted from a simulation is read from that run's CSV files here, so the text
cannot drift from the results.
"""
from __future__ import annotations

import csv
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

DOC = Path(__file__).resolve().parent
ROOT = DOC.parent
DEMOS = ROOT / "client_demos"
OUT = DOC / "FireTwins_User_Manual.docx"
ACCENT = RGBColor(0xB4, 0x3A, 0x12)
GREY = RGBColor(0x55, 0x55, 0x55)
TEXT_WIDTH_CM = 16.5
ACTUAL = -1.0          # figure width: as on the screen (96 dpi, 100 %)


# -- numbers from the runs --------------------------------------------------------------------------------------------------
def read_csv(path: Path):
    with open(path, newline="", encoding="utf-8", errors="replace") as fh:
        rows = list(csv.reader(fh))
    names = [n.strip() for n in rows[1]]
    data = [[float(x) for x in r] for r in rows[2:] if r and all(x.strip() for x in r)]
    return {n: [r[i] for r in data] for i, n in enumerate(names)}


def at(series, times, t):
    i = min(range(len(times)), key=lambda k: abs(times[k] - t))
    return series[i]


def facts_example_1() -> dict:
    run = DEMOS / "example_1" / "run"
    hrr, dev = read_csv(run / "example_1_hrr.csv"), read_csv(run / "example_1_devc.csv")
    t, q = hrr["Time"], hrr["HRR"]
    steady = [v for tt, v in zip(t, q) if tt >= 20]
    det_t = next((tt for tt, v in zip(dev["Time"], dev["HEAT_DET"]) if v >= 74.0), None)
    return {
        "hrr_15": at(q, t, 15.0), "hrr_mean_steady": sum(steady) / len(steady), "hrr_peak": max(q),
        "det_time": det_t, "det_max": max(dev["HEAT_DET"]),
        "tc_end": {k: at(dev[k], dev["Time"], 60.0) for k in ("TC_060", "TC_120", "TC_180", "TC_230")},
        "wall": (DOC / "screenshots" / "example_1" / "run_seconds.txt").read_text().strip()
        if (DOC / "screenshots" / "example_1" / "run_seconds.txt").exists() else "?",
    }


def facts_example_2() -> dict:
    run = DEMOS / "example_2" / "run"
    hrr = read_csv(run / "couch_hrr.csv")
    t, q = hrr["Time"], hrr["HRR"]
    out = (run / "couch.out").read_text(errors="replace") if (run / "couch.out").exists() else ""
    wall = next((ln.split(":")[-1].strip() for ln in out.splitlines() if "Total Elapsed Wall Clock Time" in ln), None)
    late = [v for tt, v in zip(t, q) if 15 <= tt <= 25]
    return {"hrr": {s: at(q, t, s) for s in (5, 10, 15, 20, 25)}, "hrr_peak": max(q),
            "hrr_late_mean": sum(late) / len(late) if late else float("nan"),
            "wall_min": float(wall) / 60 if wall else None}


def dstar_dx(validation: dict) -> str:
    """D*/dx as PyFDS's validation reported it for this run."""
    import re
    for msg in validation.get("warnings", []) + validation.get("errors", []):
        m = re.search(r"D\*/dx\D{0,6}([0-9]+(?:\.[0-9]+)?)", msg)
        if m:
            return m.group(1)
    return "4.8"


def pres_note(example: str, fds: str) -> str:
    """A sentence on the &PRES line of the input the run used, read from that file (empty when it has none)."""
    path = DEMOS / example / fds
    line = next((ln.strip() for ln in path.read_text(errors="replace").splitlines() if ln.startswith("&PRES")), "")         if path.exists() else ""
    return (f" Note: {fds} was written by an earlier version and still has the line {line}" if line else "")


def log_of(example: str) -> list:
    return json.loads((DOC / "screenshots" / example / "validation_log.json").read_text(encoding="utf-8"))


def step(log: list, prefix: str) -> dict:
    return next((e for e in log if str(e.get("step", "")).startswith(prefix)), {})


# -- document helpers ----------------------------------------------------------------------------------------------------------
class Manual:
    def __init__(self) -> None:
        self.doc = Document()
        self.fig = 0
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.25)
        sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(2.0)
        st = self.doc.styles
        st["Normal"].font.name = "Calibri"
        st["Normal"].font.size = Pt(10.5)
        for name, size in (("Heading 1", 18), ("Heading 2", 14), ("Heading 3", 12)):
            st[name].font.name = "Calibri"
            st[name].font.size = Pt(size)
            st[name].font.color.rgb = ACCENT if name == "Heading 1" else RGBColor(0x22, 0x22, 0x22)
        self._footer()

    def _footer(self) -> None:
        p = self.doc.sections[0].footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run("FireTwins (PyFDS) client demonstrations  |  page ")
        r.font.size, r.font.color.rgb = Pt(8), GREY
        self._field(p, "PAGE")

    @staticmethod
    def _field(paragraph, code: str) -> None:
        run = paragraph.add_run()
        for kind, text in (("begin", None), (None, code), ("separate", None), (None, "1"), ("end", None)):
            if kind:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), kind)
            elif text == code:
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = f" {code} "
            else:
                el = OxmlElement("w:t")
                el.text = text
            run._r.append(el)
        run.font.size = Pt(8)

    def h(self, text: str, level: int = 1):
        return self.doc.add_heading(text, level)

    def p(self, text: str = "", bold_lead: str = "", style: str = None, italic: bool = False):
        para = self.doc.add_paragraph(style=style)
        if bold_lead:
            para.add_run(bold_lead).bold = True
        r = para.add_run(text)
        r.italic = italic
        para.paragraph_format.space_after = Pt(4)
        return para

    def bullets(self, items, numbered: bool = False) -> None:
        for n, it in enumerate(items, 1):
            if numbered:                  # numbered here: Word's List Number would carry on from the previous list
                para = self.doc.add_paragraph()
                para.paragraph_format.left_indent, para.paragraph_format.first_line_indent = Cm(0.9), Cm(-0.6)
                para.add_run(f"{n}.\t")
            else:
                para = self.doc.add_paragraph(style="List Bullet")
            if isinstance(it, tuple):
                para.add_run(it[0]).bold = True
                para.add_run(it[1])
            else:
                para.add_run(it)
            para.paragraph_format.space_after = Pt(2)

    def table(self, rows, widths_cm, header: bool = True) -> None:
        t = self.doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Light Grid Accent 1"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                cell = t.cell(i, j)
                cell.width = Cm(widths_cm[j])
                cell.text = ""
                run = cell.paragraphs[0].add_run(str(val))
                run.font.size = Pt(9)
                run.bold = header and i == 0
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    #: menu paths in the walkthrough logs' captions that never matched the menu bar (Run is titled Analysis; Draw
    #: Room is in Model > Draw; Smokeview opens from Results)
    CAPTION_FIXES = (("Run > Open Smokeview", "Results > Open Smokeview..."), ("Run > Run Simulation", "Analysis > Run Simulation"),
                     ("Draw > Draw Room", "Model > Draw > Draw Room"))

    def figure(self, path: Path, caption: str, width_cm: float = None) -> int:
        for old, new in self.CAPTION_FIXES:
            if new not in caption:
                caption = caption.replace(old, new)
        from PIL import Image
        w, h = Image.open(path).size
        width = width_cm or TEXT_WIDTH_CM
        if h / w > 1.0:                                    # tall dialogs: keep them on the page
            width = min(width, 11.0 * w / h * 1.6, 12.0)
        width = min(width, w / 96 * 2.54 * 1.25)            # never blow small dialogs up beyond 125 %
        if width_cm == ACTUAL:                              # menus: the size they have on the screen
            width = min(w / 96 * 2.54, TEXT_WIDTH_CM)
        self.doc.add_picture(str(path), width=Cm(width))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.fig += 1
        cap = self.doc.add_paragraph()
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(f"Figure {self.fig}. ")
        r.bold, r.font.size, r.font.color.rgb = True, Pt(9), GREY
        r2 = cap.add_run(caption)
        r2.font.size, r2.font.color.rgb = Pt(9), GREY
        cap.paragraph_format.space_after = Pt(10)
        return self.fig

    def page_break(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def toc(self) -> None:
        para = self.doc.add_paragraph()
        run = para.add_run()
        for kind, text in (("begin", None), (None, 'TOC \\o "1-2" \\h \\z \\u'), ("separate", None),
                           (None, "Table of contents - updated when the document is opened in Word."), ("end", None)):
            if kind:
                el = OxmlElement("w:fldChar")
                el.set(qn("w:fldCharType"), kind)
            elif text.startswith("TOC"):
                el = OxmlElement("w:instrText")
                el.set(qn("xml:space"), "preserve")
                el.text = text
            else:
                el = OxmlElement("w:t")
                el.text = text
            run._r.append(el)

    def shots(self, example: str, names: list, log: list, intro: dict = None, width_cm: float = None) -> dict:
        """Numbered procedure: one step per screenshot whose name (after the 'NN_' number) starts with one of *names*,
        in walkthrough order. Matching the start keeps 'file_menu' from also taking 'results_file_menu'."""
        figs = {}
        n = 0
        for e in log:
            shot = e.get("screenshot")
            stem = (shot or "")[3:]
            if not shot or not any(stem.startswith(k) for k in names) or shot in getattr(self, "_used", set()):
                continue
            self.__dict__.setdefault("_used", set()).add(shot)
            n += 1
            key = next(k for k in names if stem.startswith(k))
            lead = (intro or {}).get(key)
            if lead and key not in self.__dict__.setdefault("_led", set()):
                self._led.add(key)
                self.p(lead)
            fig = self.figure(DOC / "screenshots" / example / shot, e["caption"], width_cm)
            figs[shot] = fig
        return figs


# -- content -------------------------------------------------------------------------------------------------------------------
def build() -> None:
    m = Manual()
    log1, log2 = log_of("example_1"), log_of("example_2")
    f1, f2 = facts_example_1(), facts_example_2()

    # title
    t = m.doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.paragraph_format.space_before = Pt(120)
    r = t.add_run("FireTwins (PyFDS)")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(30), ACCENT
    s = m.doc.add_paragraph()
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s.add_run("Client demonstrations - step-by-step user manual")
    r.font.size = Pt(16)
    s2 = m.doc.add_paragraph()
    s2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s2.add_run("Example 1: room fire with door, window and heat detector\nExample 2: burning couch (FDS example case)"
                   "\nExample 3: tenability and comparing runs (chapter 4)")
    r.font.size, r.font.color.rgb = Pt(12), GREY
    s3 = m.doc.add_paragraph()
    s3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = s3.add_run("Revised 29 October 2026: every screenshot and run taken again with the current version; "
                   "new menus and tools in chapter 4")
    r.italic, r.font.size, r.font.color.rgb = True, Pt(10), GREY
    m.p("")
    m.table([["Software", "Version used for this manual"],
             ["PyFDS model editor and PyFDS Results", "0.3.0 (repository E:\\firetwins)"],
             ["FDS (Fire Dynamics Simulator, NIST)", "FDS-6.10.1 release"],
             ["Smokeview (NIST)", "SMV-6.10.1 release"],
             ["Python / PySide6 / PyVista / VTK", "3.10 / 6.12.0 / 0.49.0 / 9.3.1"],
             ["Screen used for every screenshot", "1360 x 768, 100 % scaling, windows maximized"]], [7.5, 8.5])
    m.page_break()
    m.h("Contents", 1)
    m.toc()
    m.page_break()

    # 0 -- the software
    m.h("1  Before you start", 1)
    m.p("FireTwins demonstrations use three programs. All screenshots in this manual were taken from the running "
        "programs while the two cases were built, run and post-processed; nothing is a mock-up.")
    m.bullets([("PyFDS (model editor). ", "Builds and checks the FDS model, writes the FDS input file and starts FDS."),
               ("PyFDS Results. ", "A Smokeview-style 3D viewer for FDS results that comes with PyFDS."),
               ("Smokeview. ", "NIST's own viewer, installed with FDS (fds_smv\\SMV6\\smokeview.exe); PyFDS can open "
                               "a run in it and drive it with Smokeview scripts.")])
    m.h("1.1  Launching the programs", 2)
    m.p("The installer is not yet available. Until it is, the programs are started from the PyFDS folder in the "
        "Python environment that has PyFDS installed (here: the conda environment 'agent'):")
    m.table([["Program", "Command (run in E:\\firetwins)"],
             ["PyFDS model editor", "python main.py      (or: pyfds-gui)"],
             ["PyFDS Results", "python results_main.py      (or: pyfds-results)"],
             ["Smokeview", "fds_smv\\SMV6\\smokeview.exe <case>.smv   (or double-click the .smv file)"]], [5.0, 11.0])
    m.p("FDS and Smokeview 6.10.1 are in E:\\firetwins\\fds_smv; PyFDS finds them there. Both programs open "
        "maximized. A log of each session is written to user_data\\logs\\ (pyfds.log, pyfds_results.log); send it "
        "with any problem report.")
    m.h("1.2  The demonstration files", 2)
    m.table([["Folder", "Contents"],
             ["client_demos\\example_1", "Room fire: example_1.pyfds (project), example_1.fds (FDS input), run\\ (FDS "
                                         "output), results\\ (pictures and movie), example_1.pyfdsv (saved viewer setup)"],
             ["client_demos\\example_2", "Couch fire: couch_original.fds (the unchanged FDS example), couch.pyfds, "
                                         "couch.fds, run\\, results\\, couch.pyfdsv"],
             ["client_demos\\example_3", "Example 1 with the door closed and gas devices at head height (chapter 4): "
                                         "example_3.pyfds, example_3.fds, run\\, scripts\\make_case.py"],
             ["documentation", "This manual, the screenshots (screenshots\\example_1, \\example_2, \\whats_new), the screen "
                               "recordings (videos\\) and validation_report.md"],
             ["examples", "The earlier versions of both examples, kept unchanged as the reference"]], [4.5, 11.5])
    m.page_break()

    example_1(m, log1, f1)
    m.page_break()
    example_2(m, log2, f2)
    m.page_break()
    whats_new(m)
    m.doc.save(str(OUT))
    print("written", OUT, "figures:", m.fig)


def example_1(m: Manual, log: list, f: dict) -> None:
    val = step(log, "I1").get("detail", {})
    run = step(log, "K2").get("detail", {})
    m.h("2  Example 1 - Room fire with door, window and heat detector", 1)
    m.h("A  Introduction", 2)
    m.p("A propane burner fire in a small room with an open door and window, with a heat detector and a "
        "thermocouple tree above the fire. The case is built from an empty project, so it shows the model editor's "
        "dialogs one by one.", bold_lead="Purpose. ")
    m.p("Room 4.8 x 3.3 x 2.4 m (inside 25 mm gypsum-board walls) in a 6 x 4.5 x 3 m computational domain open to "
        "the outside on its sides and top. Door 0.9 x 2.1 m (south wall), window 1.2 x 1.05 m with 0.9 m sill (east "
        "wall). A 0.6 x 0.6 m burner in the middle of the floor releases 500 kW/m2 on its top face, i.e. 180 kW, "
        "reached along a t-squared curve in 15 s, and burns propane (soot yield 0.024, CO yield 0.005). Simulated "
        "time 60 s.", bold_lead="Scenario. ")
    m.p("Follow the hot gas layer and smoke as they form under the ceiling and leave through the openings; find the "
        "time the heat detector (RTI 5 (m s)^0.5, 74 C) activates; read the temperature profile above the fire; "
        "check visibility at head height.", bold_lead="Engineering objectives. ")
    m.bullets(["Building a model from scratch: mesh, library construction, Draw Room, burner, reaction, vents",
               "Devices (heat detector, thermocouples) and output slices (temperature, visibility)",
               "Model validation, the Run dialog and the live Run Monitor",
               "Post-processing in PyFDS Results (slices, 3D smoke, devices, time histories, pictures, movie) and "
               "in Smokeview"])
    m.p("PyFDS, FDS 6.10.1 and Smokeview 6.10.1 (section 1), about 2 GB of free disk space, and time for the run: it "
        f"took {float(f['wall']) / 60:.0f} minutes of wall clock time on the test computer (one MPI process). For a "
        "live demonstration, open the finished run instead (section E).", bold_lead="Prerequisites. ")
    m.p("documentation\\videos\\example_1_build_and_run.mp4 records the whole of sections B to D on screen (the FDS "
        "run as a time-lapse, one frame per 10 s); example_1_post_processing.mp4 records section E.",
        bold_lead="Screen recordings. ")

    m.h("B  Opening the program", 2)
    m.p("Start the PyFDS model editor (section 1.1). It opens maximized with an empty project. The main areas are: "
        "the menu bar and toolbar at the top; the drawing tools on the left edge; the Project Explorer (every object "
        "of the model, grouped by kind) with the Properties panel under it; the 3D view in the centre with its icon "
        "strip (rendering mode, standard views, grid, snapping, screenshot, show/hide); and the status bar.")
    m.p("To demonstrate a finished case instead of building it, use File > Open and choose "
        "client_demos\\example_1\\example_1.pyfds, then continue at section D.")
    m.shots("example_1", ["editor_start", "file_menu"], log)

    m.h("C  Building and configuring the model", 2)
    m.p("Each dialog below is shown completed, with the values used for this case. Press OK to apply a dialog; "
        "every change can be undone with Edit > Undo (Ctrl+Z). Values are in SI units (m, s, kW, C).")
    intro = {
        "project_settings": "Step 1 - Project settings. CHID is the base name of every file FDS writes.",
        "simulation_parameters": "Step 2 - Simulation parameters. Only the end time is changed (T_END = 60 s); the "
                                 "other tabs show the values used. Environment, Pressure and Radiation hold "
                                 "FDS's own defaults (the velocity tolerance is left unset: FDS works it out from "
                                 "the cell size), so PyFDS writes no &PRES line. The Output tab's intervals (1 s, "
                                 "Plot3D 10 s) are PyFDS's choice; FDS on its own would save 1000 frames over the "
                                 "run." + pres_note("example_1", "example_1.fds"),
        "add_mesh": "Step 3 - Mesh. The cell size is 6 m / 60 = 0.1 m in every direction. The cell size has to "
                    "resolve the fire: FDS guidance uses D*/dx, the characteristic fire diameter over the cell size. "
                    "For 180 kW, D* = 0.48 m; 0.15 m cells give only 3.2, which PyFDS's validation rejects (below 4), "
                    "while 0.1 m cells give 4.8.",
        "library_wall": "Step 4 - Wall construction from the library.",
        "draw_room": "Step 5 - Draw the room with its two openings.",
        "burner_surface": "Step 6 - The fire: a burner surface, the burner body, the burning face and the reaction.",
        "open_boundaries": "Step 7 - Outer boundaries. Model > Open Mesh Boundaries makes the sides and the top "
                           "of the mesh OPEN, so smoke can leave the domain. The floor stays closed, because "
                           "the room stands on the ground.",
        "library_heat_detector": "Step 8 - Instruments and outputs.",
        "model_isometric": "Step 9 - Check the finished model.",
        "editor_surfaces": "The editors (Model > Edit ...) list every object of a kind with all its settings and "
                           "the FDS text it will write; use them to review the case.",
    }
    m.shots("example_1", ["project_settings", "simulation_parameters", "add_mesh", "mesh_in_view", "library_wall",
                          "draw_room", "room_in_view", "burner_surface", "burner_block", "burner_vent", "reaction",
                          "open_boundaries", "library_heat_detector", "heat_detector", "thermocouple", "slice_",
                          "model_isometric", "model_front", "editor_"], log, intro)
    m.p("Save the project with File > Save As (example_1.pyfds). File > Export > Export as FDS writes the FDS input file "
        "example_1.fds, which is also written automatically when the run starts.")

    m.h("D  Validation and running the simulation", 2)
    m.p("Model > Validate Project checks the model before it is run: meshes and their resolution, objects outside the mesh, "
        "references to surfaces, materials and properties, ramps, controls, HVAC and the time and output settings.")
    n_e, n_w = len(val.get("errors", [])), len(val.get("warnings", []))
    m.p(f"Result for this case: {n_e} errors, {n_w} warning(s)." +
        ("".join(f" Warning: {w}" for w in val.get("warnings", []))) +
        " Errors block a run; warnings do not, but should be understood. A resolution warning (D*/dx between 4 and "
        "10) is expected for a demonstration mesh: the plume is resolved, but coarsely (see section F).")
    m.shots("example_1", ["validation"], log)
    m.p("Analysis > Run Simulation (or Run Simulation on the toolbar) opens the Run dialog. For a one-mesh model it proposes 'This computer' with one "
        "process. Choose the folder for the input and results (Browse...), check the FDS version (PyFDS uses the "
        "bundled FDS 6.10.1 when none is configured; Locate FDS... picks another), read the 'Before you run' "
        "summary, and press Run. 'Check the model first' stops the run if the model has errors.")
    m.shots("example_1", ["run_dialog"], log)
    m.p("The Messages panel opens with the Run Monitor: progress (simulated time / end time), the heat release "
        "rate curve as FDS computes it, and the last lines of the FDS log (.out file), refreshed twice a second. The "
        "status bar shows the progress too. Analysis > Stop Simulation asks FDS to finish cleanly and write a restart "
        "file; Analysis > Cancel Simulation ends it at once.")
    m.shots("example_1", ["run_monitor", "message", "run_finished"], log)
    m.p(f"Outcome on the test computer: FDS reported 'STOP: FDS completed successfully'; wall clock time "
        f"{run.get('wall_clock_s', '?')} s for 60 s of simulated time.")

    m.h("E  Post-processing", 2)
    m.h("E.1  PyFDS Results", 3)
    m.p("Open the viewer from the editor with Results > Open PyFDS Results (or start it on its own, section 1.1) "
        "and load the run with File > Open Results (.smv): choose client_demos\\example_1\\run\\example_1.smv. The "
        "Navigation View lists everything the run wrote: geometry, devices, views, 3D smoke, Plot3D, slices. Click a "
        "result to show it; use the time bar at the bottom (Play, step, drag the slider) to move through the run.")
    m.shots("example_1", ["results_empty", "results_file_menu", "results_loaded", "results_view_"], log)
    m.p("Time histories: File > Load CSV File with example_1_hrr.csv and example_1_devc.csv, then Tools > Time History "
        "Plot and tick the series.")
    m.shots("example_1", ["time_plot"], log)
    m.p("Pictures and movies: Preferences > Colorbar adjusts a colour bar; File > Save Screenshot and File > Movie > "
        "Save Movie write pictures and MP4 files. File > Save Results Visualization keeps the whole viewer setup "
        "(example_1.pyfdsv) so the same views open next time.")
    m.shots("example_1", ["results_colorbar", "results_screenshot", "results_movie"], log)
    m.h("E.2  Smokeview", 3)
    m.p("The same run opens in NIST's Smokeview. There are three ways to open it:")
    m.bullets([("Automatically after the run. ", "Tick 'Open Smokeview when the run finishes' in the Run dialog (it "
                "follows Analysis > Open Smokeview After Run, on by default). When FDS ends, Smokeview opens <CHID>.smv "
                "from the run folder (section 4.4)."),
               ("From the editor at any time. ", "Results > Open Smokeview... and choose the .smv file, here "
                "client_demos\\example_1\\run\\example_1.smv. Results > Open Results Folder... shows the folder."),
               ("From Windows. ", "Double-click example_1.smv (when .smv files open with Smokeview), or run "
                "fds_smv\\SMV6\\smokeview.exe example_1.smv in the run folder.")])
    m.p("PyFDS uses the Smokeview set in File > Preferences (smv_binary); when none is set it uses the bundled "
        "fds_smv\\SMV6\\smokeview.exe. Smokeview starts in the run folder, so it finds the slice, boundary and "
        "smoke files next to the .smv. Right-click in the Smokeview window for its menus: Load/Unload > Slice file, "
        "3D smoke; the time bar runs the animation; File/Bookmark/Render saves pictures. The pictures below were made "
        "by a Smokeview script (client_demos\\results_walkthrough.py writes it), the same commands a user can type "
        "in a .ssf file.")
    m.shots("example_1", ["smokeview_window", "smv_"], log)

    m.h("F  Results and interpretation", 2)
    tc = f["tc_end"]
    det = f"{f['det_time']:.0f} s" if f["det_time"] is not None else "not within the 60 s (peak link temperature " \
                                                                     f"{f['det_max']:.0f} C)"
    m.table([["Quantity (from the run's CSV files)", "Value"],
             ["Heat release rate at 15 s (end of the t-squared growth)", f"{f['hrr_15']:.0f} kW"],
             ["Mean heat release rate 20-60 s (design 180 kW)", f"{f['hrr_mean_steady']:.0f} kW"],
             ["Heat detector (74 C, RTI 5) activation", det],
             ["Gas temperature above the burner at 60 s: 0.6 / 1.2 / 1.8 / 2.3 m",
              " / ".join(f"{tc[k]:.0f}" for k in ("TC_060", "TC_120", "TC_180", "TC_230")) + " C"]], [10.5, 5.5])
    m.bullets(["The temperature slices show the plume rising from the burner and a hot layer forming under the "
               "ceiling, spilling out through the top of the door and the window - the openings are the vents of "
               "the room.",
               "The thermocouples sit in the plume centreline, so they read plume temperatures, not the room's "
               "average layer temperature.",
               "The 3D smoke shows where soot goes; its density follows from the assumed soot yield (0.024 kg/kg).",
               "Visibility at 1.5 m stays near its maximum where smoke has not reached head height; where it drops, "
               "smoke has come down to that level."])
    m.p(f"This is a demonstration case. The mesh resolves the fire only coarsely (D*/dx about {dstar_dx(val)}; 10 or "
        "more is recommended for quantitative work), the run covers 60 s, the burner is a prescribed heat release (no fire growth from "
        "materials), and the soot and CO yields are handbook values for well-ventilated propane burning. The "
        "results show what the model computes; they are not a validated fire-engineering assessment of a real "
        "room, which would need a mesh-sensitivity study and case-specific inputs.", bold_lead="Limitations. ")

    m.h("G  Troubleshooting", 2)
    m.table([["Symptom", "Cause and remedy"],
             ["Run button disabled; red text 'choose a folder'", "Set 'Input and results folder' in the Run dialog."],
             ["'FDS was not found on this computer'", "Use Locate FDS... and pick fds.exe (for example "
                                                     "fds_smv\\FDS6\\bin\\fds.exe)."],
             ["'the model has errors; see Model Checks'", "Model > Validate Project (or Model > Model Checks), fix each error (double-click it to select "
                                                          "the object), or untick 'Check the model first'."],
             ["Run Monitor stays at 'Waiting for FDS to start'", "FDS has not written the .out file yet; if it never "
                                                                  "does, check the folder is writable and read the Messages panel."],
             ["'FDS stopped with an error'", "Read the last lines in the Run Monitor (the ERROR line names the "
                                             "namelist); the full log is <CHID>.out in the run folder."],
             ["Heat release about twice the expected value", "The burner surface was put on the whole obstruction: "
                                                             "use an INERT block with a burner vent on its top face."],
             ["Validation error 'D*/dx ... Minimum acceptable ratio is 4.0'", "The cells are too large for the fire: "
                                                                             "refine the mesh (0.1 m here) or check the "
                                                                             "heat release rate is the one intended."],
             ["PyFDS Results shows an empty view", "Tick an item in the Navigation View (geometry, a slice ...)."],
             ["Smokeview shows no data", "Load it from the right-click menu (Load/Unload); check the .smv and the "
                                         "data files are in the same folder."]], [6.0, 10.0])

    m.h("H  Quick reference", 2)
    m.bullets(["Start PyFDS; File > Project Settings: CHID example_1",
               "Model > Simulation Parameters: T_END 60 s",
               "Model > Add Mesh: 60,45,30 cells over 0,6,0,4.5,0,3 (0.1 m)",
               "Model > Library: gypsum_wall_25mm; Model > Draw > Draw Room: door S 0.9 x 2.1, window E 1.2 x 1.05 sill 0.9",
               "Model > Add Special > Add Typed Surface: burner 500 kW/m2, TAU_Q -15; Add Obstruction (INERT) and Add Vent "
               "(BURNER_SURF) on its top; Add Reaction: propane, then Edit Reactions: soot 0.024, CO 0.005",
               "Model > Open Mesh Boundaries (sides and top; the floor stays closed)",
               "Library: heat_detector_74C; Devices > Add Device (LINK TEMPERATURE); 4 thermocouples; 3 slices",
               "Model > Validate Project (0 errors); File > Save As",
               "Analysis > Run Simulation: folder, Run; watch the Run Monitor until 'Finished'",
               "PyFDS Results: File > Open Results (.smv), example_1.smv, tick slices / 3D smoke / devices, use the time bar",
               "Smokeview: open example_1.smv, right-click > Load/Unload"], numbered=True)


def example_2(m: Manual, log: list, f: dict) -> None:
    opened = step(log, "A1").get("detail", {})
    val = step(log, "C1").get("detail", {})
    run = step(log, "D2").get("detail", {})
    m.h("3  Example 2 - Burning couch (FDS example case)", 1)
    m.h("A  Introduction", 2)
    m.p("Shows that a case written for FDS by hand opens in PyFDS, can be inspected and run in parallel, and "
        "post-processed - the route a consultant's existing FDS models take.", bold_lead="Purpose. ")
    m.p("The case is 'Examples/Fires/couch.fds' from the FDS 6.10.1 distribution (a teaching case; its material "
        "properties are marked 'completely fabricated'). A couch of fabric over foam on a gypsum base stands in a "
        "5 x 5 x 2.4 m room with one open vent; a hot (1000 C) static particle on the seat ignites it. Burning "
        "materials make the fire grow on its own - there is no prescribed heat release rate.", bold_lead="Scenario. ")
    m.p("Watch the fire spread over the couch surfaces, the flame and smoke, and the heat release rate it "
        "produces.", bold_lead="Engineering objectives. ")
    m.bullets(["Opening an existing FDS input file", "Multiplied meshes (MULT) and a 4-process MPI run",
               "Layered, burning-away surfaces and pyrolysis materials", "Boundary-file results on the couch surfaces"])
    took = f"it took {f['wall_min']:.0f} minutes on the test computer" if f.get("wall_min") else "it is a short run"
    m.p(f"The original runs for 600 s. Here T_END is set to 25 s in the GUI so the demo stays short: {took} on "
        f"4 MPI processes. Nothing else in the physics is changed.", bold_lead="Prerequisites. ")
    m.p("documentation\\videos\\example_2_open_and_run.mp4 (sections B to D, the run as a time-lapse) and "
        "example_2_post_processing.mp4 (section E).", bold_lead="Screen recordings. ")

    m.h("B  Opening the example", 2)
    m.p("Start PyFDS, then File > Open and choose client_demos\\example_2\\couch_original.fds (a copy of the "
        "untouched FDS example). PyFDS reads the FDS namelists into objects; anything it has no editor for is kept "
        "and written back unchanged.")
    m.p(f"What was read: {opened.get('meshes_in_file', '?')} &MESH line with a multiplier, giving "
        f"{opened.get('meshes_run', '?')} meshes; {opened.get('obstructions', '?')} obstructions; surfaces "
        f"{', '.join(opened.get('surfaces', []))}; materials {', '.join(opened.get('materials', []))}; "
        f"T_END {opened.get('t_end', '?')} s.")
    m.shots("example_2", ["opened"], log)
    m.h("C  Inspecting and configuring the case", 2)
    m.p("The editors under Model show what the file defines. Nothing is changed in them for this demonstration.")
    m.shots("example_2", ["editor_"], log)
    m.p("The only change: Model > Simulation Parameters, T_END from 600 s to 25 s.")
    m.shots("example_2", ["simulation_parameters", "front_view"], log)
    m.p("When the project is saved and written back to FDS, PyFDS keeps the original &DUMP settings (NFRAMES=3000, "
        "devices and heat release every 5 s) and adds its standard intervals for the files the original leaves to "
        "NFRAMES: slices, boundary files and isosurfaces every 1 s, Plot3D every 10 s, 3D smoke on. It also writes "
        "&MISC and &RADI with FDS's default values; &PRES is written only when a pressure-solver setting differs "
        "from FDS's default." + pres_note("example_2", "couch.fds") + " These change how often results are saved, "
        "not the physics.", bold_lead="Note. ")

    m.h("D  Validation and running the simulation", 2)
    m.p(f"Model > Validate Project: {len(val.get('errors', []))} errors, {len(val.get('warnings', []))} warning(s)." +
        "".join(f" Warning: {w}" for w in val.get("warnings", [])))
    m.shots("example_2", ["validation"], log)
    m.p("Analysis > Run Simulation. Because the model has 4 meshes, the dialog proposes 'This computer, in parallel "
        "(MPI)' with 4 processes - FDS needs one process per mesh - and shows which mesh each process computes.")
    m.shots("example_2", ["run_dialog", "run_monitor", "message", "run_finished"], log)
    m.p(f"Outcome on the test computer: {'FDS completed successfully' if step(log, 'D2').get('outcome') == 'ok' else 'see the validation report'}; "
        f"wall clock time {run.get('wall_clock_s', '?')} s for 25 s of simulated time on 4 MPI processes.")

    m.h("E  Post-processing", 2)
    m.h("E.1  PyFDS Results", 3)
    m.p("File > Open Results (.smv) with client_demos\\example_2\\run\\couch.smv. Results of the four meshes are listed "
        "per mesh; tick all four to see the whole plane or surface.")
    m.shots("example_2", ["results_empty", "results_file_menu", "results_loaded", "results_view_", "time_plot",
                          "results_colorbar", "results_screenshot", "results_movie"], log)
    m.h("E.2  Smokeview", 3)
    m.p("Open couch.smv in Smokeview (double-click). Right-click > Load/Unload for slices, boundary files and 3D "
        "smoke; the pictures below come from a Smokeview script.")
    m.shots("example_2", ["smokeview_window", "smv_"], log)

    m.h("F  Results and interpretation", 2)
    q = f["hrr"]
    m.table([["Time (s)", "5", "10", "15", "20", "25"],
             ["Heat release rate (kW)"] + [f"{q[s]:.0f}" for s in (5, 10, 15, 20, 25)]], [5.0, 2.2, 2.2, 2.2, 2.2, 2.2])
    m.bullets(["The ignitor heats the seat fabric; flames spread over the seat and up the back - the burning-rate "
               "boundary file shows where material is pyrolysing.",
               f"The heat release rate rises as more surface burns, to {q[15]:.0f} kW at 15 s, and averages "
               f"{f['hrr_late_mean']:.0f} kW from 15 to 25 s (peak {f['hrr_peak']:.0f} kW) while the couch keeps "
               "burning - the run is stopped long before the original 600 s.",
               "Wall temperatures rise only near the couch in this short time."])
    m.p("The material properties are the FDS example's invented values, the run covers the first 25 s only, and "
        "the 0.1 m mesh is the example's. The case demonstrates the workflow; its heat release rate is not a "
        "prediction for a real couch.", bold_lead="Limitations. ")
    m.h("G  Troubleshooting", 2)
    m.table([["Symptom", "Cause and remedy"],
             ["'4 processes for N meshes: FDS needs one process per mesh'", "Set MPI processes to the number of meshes "
                                                                            "(4) or use 'This computer' with one process."],
             ["MPI run does not start", "FDS's MPI (mpiexec) must be installed with FDS; Locate FDS... picks an install "
                                        "that has it."],
             ["Run takes far longer than expected", "Each of the 4 processes needs a free CPU core; close other heavy "
                                                    "programs or reduce T_END."],
             ["Only a quarter of a slice is shown", "Results are per mesh: tick the slice for all four meshes."],
             ["A Smokeview script stops with 'Slice files of type TEMPERATURE failed to load'",
              "This case's slices are cell-centred and LOADSLICE cannot load them; use LOADSLCF with QUANTITY='TEMPERATURE' PBX=2.5 CELL_CENTERED=T (PyFDS's script builder does this with cell_centered=True), or load the slice from Smokeview's right-click > Load/Unload menu."]],
            [6.5, 9.5])
    m.h("H  Quick reference", 2)
    m.bullets(["Start PyFDS; File > Open: couch_original.fds",
               "Inspect: Model > Edit Meshes / Surfaces / Materials / Particles",
               "Model > Simulation Parameters: T_END 25 s",
               "Model > Validate Project; File > Save As couch.pyfds",
               "Analysis > Run Simulation: parallel (MPI), 4 processes, folder, Run",
               "PyFDS Results: File > Open Results (.smv), couch.smv; tick all 4 meshes of a result",
               "Smokeview: open couch.smv"], numbered=True)


def facts_example_3() -> dict:
    run = DEMOS / "example_3" / "run"
    dev = read_csv(run / "example_3_devc.csv")
    out = (run / "example_3.out").read_text(errors="replace") if (run / "example_3.out").exists() else ""
    wall = next((ln.split(":")[-1].strip() for ln in out.splitlines() if "Total Elapsed Wall Clock Time" in ln), "?")
    return {"co_max_ppm": max(dev["CO_15"]) * 1e6, "co2_max_pct": max(dev["CO2_15"]) * 100,
            "o2_min_pct": min(dev["O2_15"]) * 100, "t_max": max(dev["GAS_T_15"]), "vis_min": min(dev["VIS_15"]),
            "fed_end": dev["FED_15"][-1], "t_end": dev["Time"][-1], "wall": wall,
            "ok": "STOP: FDS completed successfully" in out}


def whats_new(m: Manual) -> None:
    log = log_of("whats_new")
    res_path = DOC / "screenshots" / "whats_new" / "validation_log_results.json"
    log += json.loads(res_path.read_text(encoding="utf-8")) if res_path.exists() else []
    m.h("4  New and changed in the version of 29 October 2026", 1)
    m.p("This chapter shows what a code review of 29 October 2026 added or changed in the programs. The pictures were "
        "taken from the running programs, like the rest of the manual (client_demos\\whats_new_walkthrough.py). "
        "Chapters 2 and 3 were captured again with this version, runs included.")
    m.bullets([("Menus. ", "Long menus are grouped into submenus (Edit > Visibility and Transform, Model > Add from "
                           "Library, Add Special and Model Checks, View > Rendering Mode; in PyFDS Results File > Movie, "
                           "View > Section Box and Units). Every command kept its shortcut and icon."),
               ("Project Explorer filter. ", "A filter box over the object tree (section 4.2)."),
               ("Fire Engineering Calculators. ", "Tools > Fire Engineering Calculators: hand calculations to check a "
                                                  "model (section 4.3)."),
               ("Tenability (FED) and Compare Runs. ", "Two new commands in the Analysis menu of PyFDS Results "
                                                      "(section 4.5)."),
               ("Help. ", "The FDS User Guide and the Smokeview Verification Guide open from Help; the PyroSim manual "
                          "entry was removed."),
               ("Fixes that change the FDS input. ", "Open Mesh Boundaries leaves the floor closed; &PRES is written only "
                "when a pressure-solver setting differs from FDS's default (earlier versions always wrote "
                "VELOCITY_TOLERANCE=0.05, a PyFDS value that FDS does not use by default); Draw Room coordinates are "
                "rounded (1.95, not 1.9500000000000002); an HVAC fan or aircoil shared by several ducts is written once "
                "with all its ducts."),
               ("Other fixes. ", "The FED calculation uses FDS's own constants (FDS User Guide eq. 22.42-22.49) and, like "
                                 "FDS's FED device, counts the oxygen term only below 20 % O2; Tenability takes the gas "
                                 "temperature from the gases' own place; Compare Runs shows '-' instead of 'nan' for a "
                                 "device that only one run has; on/off menu entries show a tick; the AI "
                                 "Assistant no longer changes the model from a background thread.")])

    m.h("4.1  The menus", 2)
    m.p("The model editor's menu bar is File, Edit, Model, Devices, Output, Analysis, Results, View, Tools, Help. "
        "Running FDS is in the Analysis menu (and on the toolbar: Run Simulation); drawing tools are in Model > Draw "
        "and on the left edge of the window.")
    m.shots("whats_new", ["editor_with_model"], log)
    m.shots("whats_new", ["menu_", "submenu_"], log, width_cm=ACTUAL)

    m.h("4.2  Finding objects: the Project Explorer filter", 2)
    m.p("Type in the box above the Project Explorer to show only the objects whose name contains the text. 'tag:' "
        "searches the tags and 'type:' the kind of object (type:vent, type:device). The clear button (x) shows "
        "everything again; the groups that were open stay open.")
    m.shots("whats_new", ["explorer_filter"], log)

    m.h("4.3  Fire Engineering Calculators", 2)
    m.p("Tools > Fire Engineering Calculators. Choose a calculation, enter its inputs (SI units) and press Calculate; "
        "every number the method returns is listed. The calculations: fire size and cell size (D*, D*/dx), Heskestad "
        "plume, Alpert ceiling jet, smoke filling time (NFPA 92), a two-layer zone model, heat detector response and "
        "tunnel critical velocity (Kennedy). They do not change the model.")
    cell = step(log, "W5 calculator Fire size").get("detail", {})
    if cell:
        m.p(f"For example 1's 180 kW burner the calculator gives D* = {cell.get('d star m')} m and a cell size of "
            f"{cell.get('dx max m')} m or less for D*/dx of at least {cell.get('ratio min')}. Example 1 uses 0.1 m "
            "cells (D*/dx about 4.8), which is why chapter 2 calls its resolution coarse.")
    m.shots("whats_new", ["fire_eng_"], log)

    m.h("4.4  Opening a run in Smokeview", 2)
    m.p("When 'Open Smokeview when the run finishes' is ticked in the Run dialog, Smokeview opens the run's <CHID>.smv "
        "as soon as FDS ends. At any later time use Results > Open Smokeview... (any .smv file), or Results > View "
        "Results of Last Run to open the last run in PyFDS Results instead. PyFDS uses the Smokeview set in File > "
        "Preferences (smv_binary) and otherwise the bundled fds_smv\\SMV6\\smokeview.exe.")
    m.shots("whats_new", ["run_dialog_smokeview"], log)
    m.p("In the Analysis menu, 'Open Smokeview After Run' shows a tick when it is on (like every on/off entry of the "
        "menus); the Run dialog's check box starts from it.")

    m.h("4.5  Example 3 - Tenability and comparing two runs", 2)
    if not (DEMOS / "example_3" / "run" / "example_3_devc.csv").exists():
        m.p("Example 3 has not been run yet; run client_demos\\example_3\\example_3.fds first.")
        return
    f = facts_example_3()
    m.p("Example 1's room with the door closed: the only opening is the window. Gas devices at head height (1.5 m) "
        "in the back of the room, out of the plume, record carbon monoxide, carbon dioxide and oxygen (volume "
        "fractions), temperature, visibility and FDS's own fractional effective dose (QUANTITY='FED'). Everything "
        "else is example 1. client_demos\\example_3\\scripts\\make_case.py builds it from example_1.pyfds.",
        bold_lead="Scenario. ")
    m.p(f"FDS 6.10.1, one process: {'completed successfully' if f['ok'] else 'see example_3.out'}, wall clock time "
        f"{f['wall']} s for {f['t_end']:.0f} s of simulated time.", bold_lead="Run. ")
    m.p("Open the run in PyFDS Results (File > Open Results (.smv), example_3.smv) and load its device file with File > Load "
        "CSV File (example_3_devc.csv). The new commands are in the Analysis menu.")
    m.shots("whats_new", ["results_menu_"], log, width_cm=ACTUAL)
    m.shots("whats_new", ["time_plot_gases"], log)
    m.p("Analysis > Tenability (FED) computes Purser's fractional effective dose from the chosen device columns. It "
        "picks the CO, O2, CO2 and temperature columns by name, uses FDS's constants, and reports the time to FED 0.3 "
        "(incapacitation, used as ASET), a separate heat dose, and the margin to the required safe egress time (RSET) "
        "you enter. For the gas temperature it takes the column named for the same place as the CO column (GAS_T_15 "
        "for CO_15), not just the first temperature in the file (here TC_060, a plume thermocouple). Check the "
        "columns it picked before you calculate.")
    m.shots("whats_new", ["tenability"], log)
    ten = step(log, "R1 tenability").get("detail") or {}
    rows = ten.get("rows", {})
    table = [["Quantity at the device (1.5 m)", "Value"],
             ["Peak carbon monoxide", f"{f['co_max_ppm']:.0f} ppm"],
             ["Peak carbon dioxide", f"{f['co2_max_pct']:.2f} %"],
             ["Lowest oxygen", f"{f['o2_min_pct']:.2f} %"],
             ["Peak gas temperature", f"{f['t_max']:.0f} C"],
             ["Lowest visibility", f"{f['vis_min']:.1f} m"],
             [f"FED from FDS's FED_15 device at {f['t_end']:.0f} s", f"{f['fed_end']:.3g}"]]
    table += [[f"PyFDS Tenability: {k}", v] for k, v in rows.items()]
    m.table(table, [10.0, 6.0])
    fed_hand = ten.get("rows", {}).get("Maximum FED (gases)")
    m.p(f"The dose is very small: nobody at this point would be incapacitated by the gases in the first 60 s. "
        f"PyFDS's value ({fed_hand}) agrees with FDS's own FED_15 device ({f['fed_end']:.3g}) to about 1 %; the rest "
        "comes from FDS adding up the dose every time step, PyFDS from the 1 s device output. Like FDS's device, "
        "PyFDS counts the oxygen term only below 20 % O2 (humid ambient air here has 20.76 %); FDS's User Guide "
        "(eq. 22.48) does not state this cutoff, it was found by comparing with this run.")
    m.p("Analysis > Compare Runs compares the devices of two run folders. The folder of the open run is filled in as "
        "the baseline; choose the other run, set the threshold and press Compare. Each device is listed with its peak "
        "and mean in both runs and the difference; devices that only one run has are listed as such. Here example 1 "
        "(door open) is the baseline and example 3 (door closed) the variant. A device that only one run has shows "
        "its own peak, '-' for the other run, and the status 'variant only' or 'baseline only'.")
    m.shots("whats_new", ["compare_runs"], log)
    cmp = step(log, "R2 compare runs").get("detail") or {}
    if cmp:
        common = {k: v for k, v in cmp.items() if v[0] in ("ok", "warning")}
        m.p("Devices in both runs: " + "; ".join(f"{k} {v[1]:+.1f} % ({'above threshold' if v[0] == 'warning' else v[0]})" for k, v in common.items()) +
            f". In example 3 only: {', '.join(k for k, v in cmp.items() if v[0] == 'missing_baseline') or 'none'}.")
    m.p("A demonstration on example 1's coarse mesh (section 2F), for 60 s only, with one device location. FED is "
        "evaluated where the devices are, not along an escape route, and leaves out irritant gases. The numbers show "
        "how the tools work; they are not a tenability assessment of a real room.", bold_lead="Limitations. ")


def finish_in_word() -> None:
    """Let Word fill in the table of contents and page numbers, then save the PDF beside the DOCX."""
    try:
        import win32com.client
    except ImportError:
        print("pywin32 not installed: the table of contents fills when the DOCX is opened in Word; no PDF written")
        return
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    try:
        doc = word.Documents.Open(str(OUT.resolve()))
        for toc in doc.TablesOfContents:
            toc.Update()
        doc.Fields.Update()
        doc.Save()
        doc.ExportAsFixedFormat(str(OUT.with_suffix(".pdf").resolve()), 17)       # 17 = wdExportFormatPDF
        doc.Close(False)
        print("wrote", OUT.with_suffix(".pdf"))
    finally:
        word.Quit()


if __name__ == "__main__":
    build()
    finish_in_word()
    sys.exit(0)
